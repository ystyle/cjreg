#!/usr/bin/env python3
"""生成决赛路演 PPT（12 页，16:9，可编辑文本与表格）。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
DARK = RGBColor(0x1A, 0x20, 0x2C)
GREY = RGBColor(0x5A, 0x63, 0x70)
LIGHT = RGBColor(0xF2, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, text="", size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def bullets(slide, x, y, w, h, items, size=17, color=DARK, space=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = it
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return tb


def title_bar(slide, title, kicker=""):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    box(slide, 0.55, 0.22, 11.5, 0.6, title, size=26, color=WHITE, bold=True)
    if kicker:
        box(slide, 0.6, 1.15, 12, 0.4, kicker, size=13, color=GREY)


def table(slide, x, y, w, rows, col_w=None, header=True, size=14):
    nrows, ncols = len(rows), len(rows[0])
    h = 0.42 * nrows
    shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            if header and ri == 0:
                r.font.bold = True
                r.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = ACCENT
            else:
                r.font.color.rgb = DARK
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if ri % 2 else WHITE
    return shape


# ---------- 1 封面 ----------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = ACCENT; bg.line.fill.background()
box(s, 0.9, 2.0, 11.5, 1.0, "cjreg · 仓颉私有中心仓与多仓体系", size=40, color=WHITE, bold=True)
box(s, 0.95, 3.05, 11.5, 0.6, "仓颉生态创新开发挑战赛 · 决赛路演", size=22, color=WHITE)
box(s, 0.95, 4.35, 11.5, 0.5, "技术课题 ｜ 版本 v0.3.0（2026-10-20）", size=16, color=WHITE)
box(s, 0.95, 4.85, 11.5, 0.5,
    "https://atomgit.com/ystyle/cjreg    ghcr.io/ystyle/cjreg:v0.3.0    https://ystyle.top/cjreg/",
    size=14, color=WHITE)

# ---------- 2 问题与定位 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "问题：官方中心仓很好，但它不解决这三件事")
bullets(s, 0.8, 1.7, 11.8, 4.6, [
    "① 内网 / 离线环境没有可用仓 —— 中心仓在公网，流水线与机器在内网",
    "② 发布不可逆，同步不可控 —— 版本不能删除、不能覆盖；索引同步从几十分钟到几天，失败要等",
    "③ 团队需要自己的仓与预演位 —— 预发布、内部包、灰度版本不该进公共仓",
    "",
    "cjreg 的定位：与官方协议互通的私有中心仓 + 多仓体系（不是替代，是补位）",
], size=19)
box(s, 0.8, 6.1, 11.8, 0.9,
    "生产实例已在内网常驻运行，承载团队日常发布、门户与审计", size=15, color=GREY)

# ---------- 3 能力地图 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "能力地图：一条发布链路上的七件事")
rows = [
    ["环节", "能力", "关键实现"],
    ["发布", "官方二进制协议互通；流式接收；幂等/冲突/越权判定", "分段解析 + 边读边算 sha256 + 原子落盘"],
    ["索引 / 下载", "NDJSON 索引；流式下载；下载计数", "IndexService / BlobStore.stream"],
    ["多仓", "官方 + 私有 + 团队仓；回源、缓存、解析诊断、重拉", "ProxyService / ResolveService / UpstreamProbe"],
    ["发布计划", "选包自动补依赖（拓扑序）→ 版本对照 → 推送；六态 + SSE", "PublishPlanService / PublishExecutor"],
    ["权限", "首个发布者 = owner；团队 write/overwrite；管理员直通", "effectivePermission 双栈裁决"],
    ["门户 / 管理端", "用户门户（Token、我的包、我的团队）+ 管理后台", "cjxt 声明式页面 + WS 补丁"],
    ["审计 / 公开 API", "发布/认证/管理审计（IP/UA）；公开只读 API 供 CI 集成", "LogDoc / /api/stats /api/packages"],
]
table(s, 0.55, 1.45, 12.2, rows, col_w=[1.7, 5.3, 5.2], size=13)

# ---------- 4 架构 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "技术方案：分层与关键取舍", "纯仓颉实现，静态单二进制；依赖 bstorm（KV）、cjxt（UI/WS）、tang（HTTP）")
bullets(s, 0.7, 1.75, 6.0, 4.8, [
    "protocol —— 官方发布/下载二进制协议、URL 分片",
    "index —— NDJSON 生成与兼容解析",
    "store —— bstorm 文档层（包/用户/上游/团队/审计）",
    "server —— 发布裁决、代理回源、计划执行、handler",
    "ui —— cjxt 页面（公共面 / 管理面分包子包）",
    "protocol 与 handler 之间全部走 DTO 序列化",
], size=16, space=10)
bullets(s, 7.0, 1.75, 5.6, 4.8, [
    "为什么不用 SQLite：LSM + 纯仓颉、无 C 依赖，单二进制",
    "为什么自研 tar/gzip 解压：发布/回源要提取 README 与 cjpm.toml",
    "为什么 UI 用 WS 补丁：声明式信号派生，前端零手写 DOM",
    "发布链路全程流式：内存与包体解耦（决赛核心改造）",
], size=16, color=GREY, space=10)

# ---------- 5 流式发布 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "核心难点 ①：流式发布 / 下载 / 回源（内存 O(1)）")
rows = [
    ["指标", "初赛 v0.2.0", "决赛 v0.3.0"],
    ["最大可发布包体", "≤32 MiB 稳定，48–50 MiB 偶发 OOM", "200 MiB 稳定（上限 500 MiB 可配）"],
    ["200 MiB 发布耗时", "失败", "0.618 s（324 MiB/s）"],
    ["发布峰值内存 VmHWM", "≈ 包体 2–3 倍", "恒定 291 MiB（2 / 32 / 200 MiB 一致）"],
    ["200 MiB 下载", "Out of memory", "成功，sha256 一致"],
]
table(s, 0.55, 1.5, 12.2, rows, col_w=[3.4, 4.4, 4.4], size=14)
box(s, 0.6, 4.35, 12.0, 2.2,
    "做法：按段解析头 → 边读边算 sha256 → 直接写临时文件 → 校验后 rename 原子落盘；\n"
    "下载侧 BlobStream → handler 直接 body(stream)；回源同样流式落盘。\n"
    "复现：bash scripts/bench.sh（桩上游 + k6，VmHWM 取 /proc/<pid>/status）",
    size=15, color=DARK)

# ---------- 6 权限 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "核心难点 ②：权限双栈与「首个发布者 = owner」")
rows = [
    ["操作者", "发新版本", "覆盖已有版本", "说明"],
    ["包 owner", "✓", "✗（需 overwrite）", "owner = 该包最早版本发布者，唯一且不随覆盖转移"],
    ["团队成员（write）", "✓", "✗", "按团队关联的组织 / 包判定"],
    ["团队成员（overwrite）", "✓", "✓", "覆盖需显式授权"],
    ["平台管理员", "✓", "✓", "直通 overwrite（预演必需：同版本可反复发）"],
]
table(s, 0.55, 1.5, 12.2, rows, col_w=[3.0, 2.1, 2.6, 4.5], size=14)
bullets(s, 0.7, 4.55, 12.0, 2.0, [
    "退出团队后不可见是有意的产品决策：私有仓的可见性跟随权限，不做历史遗留读取",
    "越权返回 403，并在审计里记 failed + 原因（运营可对照两条记录）",
], size=15, space=10)

# ---------- 7 六态 + SSE ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "核心难点 ③：发布计划六态 + SSE 实时进度")
bullets(s, 0.7, 1.7, 12.0, 3.2, [
    "六态：pending → running → waiting_index（等索引）→ completed / failed；另有 skipped（手动跳过）",
    "「等索引」判据：版本出现、且 sha256 一致，避免索引先到、制品后到造成假完成",
    "SSE：GET /api/admin/publish-plans/:id/events（event: progress / done），轮询接口保留兜底",
    "页面实时滚动；可手动跳过某项；中断后计划状态可恢复（保持 pending 而不是假 running）",
], size=18, space=12)
box(s, 0.7, 5.3, 12.0, 1.5,
    "为什么需要它：发布不可逆 → 推送目标仓时「发没发成功」必须可见、可核验、可跳过，"
    "否则一次网络抖动就让人不敢继续", size=15, color=GREY)

# ---------- 8 迭代成果 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "决赛阶段迭代成果（相比初赛 v0.2.0）")
rows = [
    ["维度", "初赛 v0.2.0", "决赛 v0.3.0", "变化"],
    ["单测用例", "154", "1002", "×6.5"],
    ["行覆盖率", "39.3%", "91.3%", "+52.0pp"],
    ["发布计划", "四态 + 轮询", "六态 + SSE + 跳过", "新能力"],
    ["团队管理", "仅管理员", "owner 自助（成员/转移/关联）", "新能力"],
    ["交付", "宿主构建", "自包含多阶段镜像 Dockerfile.build", "新能力"],
    ["工程结构", "admin_handler 1358 行", "122 行；UI 分子包、一页一文件；DTO 化", "可维护性"],
]
table(s, 0.55, 1.45, 12.2, rows, col_w=[2.6, 3.4, 4.4, 1.8], size=13)

# ---------- 9 质量与性能证据 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "质量与性能：每个数字都可复现")
bullets(s, 0.7, 1.65, 6.1, 4.8, [
    "1002 个用例：逻辑单测 + 进程内 HTTP 夹具 + 页面渲染",
    "91.3% 行覆盖（cjcov，HTML 报告进 CI artifact）",
    "双仓 e2e 9 步：发布 6 包 → 拓扑推送 → 公开 API → 上游连通 → 三级删除 → 组织 CRUD",
    "CI：构建 + 单测 + 覆盖率 + e2e + 文档站 + 容器冒烟",
], size=16, space=10)
bullets(s, 7.0, 1.65, 5.6, 4.8, [
    "索引缓存命中：16.16 ms → 2.74 ms（5.9×，桩计数证明零回源）",
    "发布吞吐：74 / 278 / 324 MiB/s（2 / 32 / 200 MiB）",
    "并发：1/4/16/64 VU → 88 / 294 / 414 / 356 req/s（拐点 16 VU）",
    "脚本：scripts/bench.sh + bench-stub-upstream.py + k6",
], size=16, space=10)

# ---------- 10 实用价值 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "实用价值与落地")
bullets(s, 0.8, 1.7, 12.0, 4.6, [
    "内网常驻私有仓：团队日常发布、门户取 Token、管理后台维护团队与上游",
    "预演仓：发布不可逆 → 先在私有仓预演，同一版本在管理员直通下可反复发",
    "多仓体系：官方 / 私有 / 演示空站并存，解析诊断直接回答「该发哪个仓」",
    "CI 集成：公开只读 API（/api/stats、/api/packages...）与 cjpm 客户端兼容",
    "一键部署：docker compose + 数据卷 + 优雅关闭；镜像 ghcr.io/ystyle/cjreg:v0.3.0",
], size=19, space=14)

# ---------- 11 限制与后续 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "如实说明：未做的与后续路线")
rows = [
    ["项", "现状", "计划"],
    ["多架构镜像", "仅 linux/amd64", "arm64（buildx + QEMU）"],
    ["镜像体积", "507 MB（archlinux + 运行时）", "换底座/瘦身，保持两阶段 glibc 一致"],
    ["镜像预热 / 周期同步", "未做（首次回源有冷启动）", "serve --sync-interval"],
    ["K8s 清单", "未做（提供 docker compose）", "可选补充"],
]
table(s, 0.55, 1.5, 12.2, rows, col_w=[3.6, 4.4, 4.2], size=14)
box(s, 0.6, 4.5, 12.0, 1.9,
    "取舍原则：先把「发布/索引/多仓/权限」主链路的正确性与可复现性做扎实，"
    "再把交付形态（多架构、体积、编排）补齐；功能缺失与限制全部写在仓库 docs 里。",
    size=15, color=GREY)

# ---------- 12 收尾 ----------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = ACCENT; bg.line.fill.background()
box(s, 0.9, 2.2, 11.5, 0.9, "谢谢 · Q & A", size=40, color=WHITE, bold=True)
box(s, 0.95, 3.5, 11.5, 1.8,
    "仓库  https://atomgit.com/ystyle/cjreg  （tag v0.3.0）\n"
    "镜像  ghcr.io/ystyle/cjreg:v0.3.0\n"
    "文档  https://ystyle.top/cjreg/   验证 docs/verification.md   基准 docs/benchmark.md",
    size=16, color=WHITE)


# ---------- 演讲者备注（演讲者视图直接用；每页 30–60 秒）----------
NOTES = [
    "【第1页 · 约 30 秒】开场：我是 XXX，作品是 cjreg —— 仓颉私有中心仓与多仓体系，技术课题。"
    "一句话：官方中心仓解决公共分发，我们解决内网/离线/预演这三件它不覆盖的事。",
    "【第2页 · 约 60 秒】三条各说一句就够：内网没仓、发布不可逆+同步不可控、团队要自己的预演位。"
    "落点：不是替代官方仓，是补位，协议完全互通。",
    "【第3页 · 约 60 秒】不要逐行念表格。挑三件讲：发布（协议互通+流式）、多仓（回源+解析诊断）、"
    "发布计划（拓扑推送+实时进度）。其余一句『都在同一个仓里』带过。",
    "【第4页 · 约 50 秒】分两列讲：左边是分层（protocol/index/store/server/ui），右边是取舍。"
    "必讲的一句：『发布链路全程流式，这是决赛最大的改造』。",
    "【第5页 · 约 70 秒】先给对比表里最刺眼的一行：200 MiB 从失败到 0.618 秒、内存恒定 291 MiB。"
    "再用三句话说做法：边读边算 sha256、写临时文件、rename 原子落盘。最后说复现方式（bench.sh）。",
    "【第6页 · 约 40 秒】讲矩阵的两端：owner 能发新版本但不能覆盖；管理员直通 overwrite（预演必需）。"
    "退出团队即不可见是**有意的产品决策**，不是缺陷。",
    "【第7页 · 约 50 秒】一句话说为什么需要：发布不可逆，所以『发没发成功』必须可见可核验。"
    "判据要点：版本出现且 sha256 一致才算完成，避免索引先到的假完成。",
    "【第8页 · 约 60 秒】只念五个数字：154→1002、39.3%→91.3%、六态+SSE、owner 自助、自包含构建。"
    "每个数字都能在仓库里找到脚本和报告。",
    "【第9页 · 约 40 秒】左列说『测试怎么组织』（逻辑单测 + 进程内 HTTP 夹具 + 页面渲染），"
    "右列只念两个数字：索引 5.9×、并发拐点 16 VU / 414 req/s。",
    "【第10页 · 约 40 秒】四个场景各一句：内网常驻、预演仓、多仓解析、CI 集成。"
    "补一句生产在跑、有真实团队与审计数据。",
    "【第11页 · 约 40 秒】主动说没做的三件（arm64、瘦身、预热），再给取舍原则——"
    "先主链路正确性，再交付形态。主动交代比被问出来得分高。",
    "【第12页 · 约 50 秒】谢谢 + 给链接（仓库 tag v0.3.0 / 镜像 / 文档站）。"
    "留一句钩子：『发布链路的流式改造和测试体系是这次迭代最重的两块，评委老师想先听哪块？』",
]

for _slide, _note in zip(prs.slides, NOTES):
    _slide.notes_slide.notes_text_frame.text = _note

out = "/home/ystyle/Projects/Cangjie/cjreg/docs/finals-roadshow-v0.3.0.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides))
